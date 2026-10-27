"""Build results/Lang-1-2-3_experiments.xlsx from results/experiments/*.csv (run scripts/run/run_experiments.py and
scripts/run/cosine_audit.py first). One sheet per results table, a formula-driven Summary, charts, a bug table,
the cosine audit, the literature table and a metric guide.
Usage: envs\\modern312\\Scripts\\python scripts\\run\\build_excel_experiments.py"""
import json
import os
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
sys.path.insert(0, HERE)
import build_excel as BE  # noqa: E402  (shared styles/helpers)

EXP = os.path.join(ROOT, "results", "experiments")
OUT = os.path.join(ROOT, "results", "Lang-1-2-3_experiments.xlsx")
F, FB, FS, FI, FT = BE.F, BE.FB, BE.FS, BE.FI, BE.FT


def csv(name):
    return pd.read_csv(os.path.join(EXP, name))


def col_letter(df, name):
    return get_column_letter(list(df.columns).index(name) + 1)


def data_sheet(wb, title, df, widths=None, freeze="B2", numfmt=None):
    ws = wb.create_sheet(title)
    BE.write_table(ws, df)
    ws.freeze_panes = freeze
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(df.columns)), len(df) + 1)
    BE.autowidth(ws, maxw=48)
    ws.row_dimensions[1].height = 42
    return ws


def bugs_table():
    rows = []
    info = {1: ("LANG-747", "NumberUtils does not handle Long Hex numbers", "wrong condition (counts hex digits instead of checking the value)",
                "d1a45e9738de5b3e299bb51e987565dcce55fee6", "math/NumberUtils.java", "467, 468, 471", 173, 981),
            2: ("LANG-879", "LocaleUtils test fails with new Locale \"ja_JP_JP_#u-ca-japanese\" of JDK7", "missing validation (omission: guard clause absent)",
                None, "LocaleUtils.java", "92 (first line after the missing guard)", 14, 208),
            3: ("LANG-693", "Method createNumber from NumberUtils doesn't work for floating point numbers other than Float",
                "missing precision branches (omission of conditional guards)", None, "math/NumberUtils.java", "593, 597, 601, 605", 172, 979)}
    for n, (issue, summ, kind, fixed, bfile, lines, tests, stmts) in info.items():
        if n > 1:
            m = json.load(open(os.path.join(ROOT, "data", "bugs", "Lang-%d" % n, "manifest.json")))
            fixed = m["fixed_commit"]
        rows.append({"bug": "Lang-%d" % n, "jira issue": issue, "bug report summary": summ, "bug type": kind,
                     "fixed commit (apache/commons-lang)": fixed, "buggy file": bfile, "buggy lines (answer key)": lines,
                     "tests run (Defects4J relevant)": tests, "statements in spectrum": stmts, "source files": 108,
                     "report has stack trace": "yes (49 frames, foreign project)" if n == 2 else "no"})
    return pd.DataFrame(rows)


LIT = [
    ("Zhou, Zhang, Lo - BugLocator, ICSE 2012", "https://ink.library.smu.edu.sg/sis_research/1531", "rVSM: tf-idf variant with a file-size correction; plus similar-bug history", "ir_models.rvsm (history part not implemented)"),
    ("Saha et al. - BLUiR, ASE 2013", "https://ischool.utexas.edu/ml/papers/saha-ase13.pdf", "Structured IR over class / method / variable names and comments", "ir_models.bluir (regex field extraction)"),
    ("Feature-location study (snippet only)", "https://static.aminer.org/pdf/20170130/pdfs/sigsoft/enipyx1fcu9xh8dzrtcz50jbq4mgskj7.pdf", "BM25 / BM25F beat VSM, unigram model and LDA on four projects", "ir_models.bm25, lm_dirichlet, lsi"),
    ("Purdue SCAM 2012 - stemming and source-code retrieval", "https://engineering.purdue.edu/RVL/Publications/SCAM12.pdf", "Benefit of stemming depends on other factors (e.g. tf-idf filtering of common terms)", "stem option, evaluated with and without"),
    ("Le, Oentaryo, Lo - IR + spectrum 'Better Together', FSE 2015", "https://cs.unibg.it/esecfse_proceedings/fse15/p579-le.pdf", "Multi-modal localisation with bug reports and spectra; large gains on 157 bugs", "motivates hybrid.py (ours is much simpler)"),
    ("Dao, Zhang, Meng - execution information for IRBL, 2017", "https://experts.illinois.edu/en/publications/how-does-execution-information-help-with-information-retrieval-ba/", "How execution information helps IR-based localisation", "background"),
    ("Sarhan, Vancsics, Beszedes - tie-breaking by call frequency", "https://arxiv.org/pdf/2108.11031", "54-56% of buggy methods in Defects4J tie with others (avg group about 6)", "tie-breaking idea; ours uses bug-report text"),
    ("Survey: DL meets IR-based bug localization (2025)", "https://arxiv.org/pdf/2505.00144", "61 studies; MAP (59), MRR, Top-k are the usual metrics", "MAP, MRR, Top@N, nDCG, recall@k"),
    ("FlexFL (2024)", "https://arxiv.org/pdf/2411.10714", "Top-N (1,3,5) closest to developer expectation; EXAM definition", "TOP@1/3/5, EXAM"),
    ("Wong et al. / SBFL survey (arXiv 1607.04347)", "https://arxiv.org/pdf/1607.04347", "Families of spectrum-based formulas", "sfl.py: 11 formulas"),
    ("Zou et al. - FL families and combinations (arXiv 1803.09939)", "https://arxiv.org/pdf/1803.09939", "Compares fault-localisation families and their combinations", "background"),
]
GUIDE = [
    ("Top@N (file)", "1 if the buggy file is in the top N files"), ("MRR", "mean of 1 / rank of the best buggy item"),
    ("MAP / AP", "mean over buggy items of (i / rank_i); unranked buggy items count 0"),
    ("nDCG@10", "discounted gain with binary relevance, normalised by the ideal ranking"),
    ("recall@k% files", "share of the buggy files inside the top k% of files (this is what top-k filtering needs)"),
    ("rank worst / avg / best", "tied statements take the last / middle / first position of their tie group"),
    ("TOP@N worst / avg", "bugs whose best buggy statement is within N (ties pessimistic / averaged)"),
    ("EXAM", "best buggy rank / total statements in the full spectrum (same denominator for every method)"),
    ("wasted effort", "statements read before the first buggy one (rank - 1)"),
    ("W/D/L vs SFL:ochiai", "bugs where the method ranks the bug better / equal / worse than Ochiai (by average rank)"),
    ("MRR CI95", "percentile bootstrap interval of the mean - very wide with 3 bugs"),
    ("tie_file / tie_text / tie_file_text", "lexicographic keys (SFL score, IR file score, statement text score): the later keys only reorder EXACT ties"),
    ("linear a b", "(1-a-b)*SFL + a*file score + b*text score, each min-max normalised"),
    ("+clean", "stack-frame lines removed from the bug report before retrieval"), ("+name", "bonus when the file's class name occurs as a word in the report"),
]


def summary(wb, ir, st):
    ws = wb.create_sheet("Summary", 0)
    ws["A1"] = "ISBL improvements on Defects4J Lang-1, Lang-2, Lang-3"
    ws["A1"].font = FT
    ws["A2"] = ("Three bugs only: results are indicative, not statistically meaningful. Nothing was tuned on these bugs; all configurations were "
                "fixed in advance. Every number below is a formula reading the data sheets.")
    ws["A2"].font = FI
    # ---- file-level table
    ws["A4"] = "A. IR step: how well each configuration ranks the buggy FILE (108 files per bug)"
    ws["A4"].font = FS
    heads = ["IR configuration", "buggy-file rank per bug (Lang-1/2/3)", "Top@1", "Top@5", "MRR", "MAP", "nDCG@10", "recall@10% files"]
    keys = [None, "ranks_per_bug", "Top@1", "Top@5", "MRR", "MAP", "nDCG@10", "recall@10%files"]
    for j, h in enumerate(heads):
        ws.cell(row=5, column=1 + j, value=h)
    BE.style_header(ws, 5, 1, len(heads))
    ws.row_dimensions[5].height = 42
    picks = ["tf_cosine | paper", "tf_cosine | paper +clean", "tfidf_authors | none", "tfidf_sublinear | paper", "tfidf_sublinear | split+stem +clean",
             "rvsm | paper", "bm25 | paper", "bm25 | split+stem +clean", "bluir_like | split+stem +clean"]
    n_ir = len(ir) + 1
    for i, name in enumerate(picks, start=6):
        ws.cell(row=i, column=1, value=name).font = FB
        for j, k in enumerate(keys[1:], start=2):
            c = col_letter(ir, k)
            cell = ws.cell(row=i, column=j, value='=INDEX(IR_File_Summary!$%s$2:$%s$%d,MATCH($A%d,IR_File_Summary!$A$2:$A$%d,0))' % (c, c, n_ir, i, n_ir))
            cell.number_format = "0.000" if k in ("MRR", "MAP", "nDCG@10", "recall@10%files") else "0"
            cell.font = F
        for c in range(1, len(heads) + 1):
            ws.cell(row=i, column=c).border = BE.BOX
    r = 6 + len(picks) + 1
    # ---- statement-level table
    ws.cell(row=r, column=1, value="B. Statement level: ranking the buggy STATEMENT (k = 50% where a filter is used)").font = FS
    heads2 = ["method", "best buggy rank per bug, ties averaged (Lang-1/2/3)", "TOP@1/3/5 worst-case ties", "MRR worst", "MRR avg",
              "MAP avg", "EXAM avg", "wins vs Ochiai", "draws", "losses"]
    ws_h = r + 1
    for j, h in enumerate(heads2):
        ws.cell(row=ws_h, column=1 + j, value=h)
    BE.style_header(ws, ws_h, 1, len(heads2))
    ws.row_dimensions[ws_h].height = 42
    n_st = len(st) + 1
    mpicks = ["SFL:ochiai", "SFL:dstar", "SFL:tarantula",
              "ISBL[tf_cosine | paper,k=0.5]+ochiai", "ISBL[tfidf_sublinear | split+stem +clean,k=0.5]+ochiai",
              "HYBRID[tfidf_sublinear | split+stem +clean,k=0.5,tie_file]+ochiai",
              "HYBRID[tfidf_sublinear | split+stem +clean,k=0.5,tie_text]+ochiai",
              "HYBRID[tfidf_sublinear | split+stem +clean,k=0.5,tie_file_text]+ochiai",
              "HYBRID[tfidf_sublinear | split+stem +clean,k=0.5,linear a=0.3]+ochiai",
              "HYBRID[tfidf_sublinear | split+stem +clean,k=0.5,linear a=0.3 b=0.3]+ochiai"]
    scol = {k: col_letter(st, k) for k in ["ranks_avg_per_bug", "TOP@1_worst", "TOP@3_worst", "TOP@5_worst", "MRR_worst", "MRR_avg", "MAP_avg",
                                           "EXAM_avg", "vs_SFL:ochiai_win", "vs_SFL:ochiai_draw", "vs_SFL:ochiai_lose"]}

    def look(i, k):
        c = scol[k]
        return 'INDEX(Statement_Summary!$%s$2:$%s$%d,MATCH($A%d,Statement_Summary!$A$2:$A$%d,0))' % (c, c, n_st, i, n_st)

    for i, name in enumerate(mpicks, start=ws_h + 1):
        ws.cell(row=i, column=1, value=name).font = FB
        ws.cell(row=i, column=2, value="=" + look(i, "ranks_avg_per_bug")).font = F
        ws.cell(row=i, column=3, value='=%s&"/"&%s&"/"&%s' % (look(i, "TOP@1_worst"), look(i, "TOP@3_worst"), look(i, "TOP@5_worst"))).font = F
        for j, (k, nf) in enumerate([("MRR_worst", "0.000"), ("MRR_avg", "0.000"), ("MAP_avg", "0.000"), ("EXAM_avg", "0.0000"),
                                     ("vs_SFL:ochiai_win", "0"), ("vs_SFL:ochiai_draw", "0"), ("vs_SFL:ochiai_lose", "0")], start=4):
            c = ws.cell(row=i, column=j, value="=" + look(i, k))
            c.number_format, c.font = nf, F
        for c in range(1, len(heads2) + 1):
            ws.cell(row=i, column=c).border = BE.BOX
    r = ws_h + len(mpicks) + 2
    ws.cell(row=r, column=1, value="Key findings").font = FS
    finds = [
        "1. The cosine arithmetic is correct (equals scikit-learn and SciPy to 1e-16). The paper's raw term-frequency weighting is the weak part: the wrong top-ranked file gets 88-100% of its score from words found in most files.",
        "2. Any idf-weighted model (tf-idf, rVSM, LSI) ranks the buggy file 1st on all three bugs; raw-tf cosine ranks it 8th, 23rd and 16th.",
        "3. Lang-2's report holds a pasted stack trace from an unrelated project; stripping stack frames (+clean) lifts the buggy file on Lang-2 for most models.",
        "4. False exclusion: at k = 10% the paper's raw-tf IR throws the buggy file away on 2 of 3 bugs; tf-idf, rVSM and BLUiR-style IR keep it while cutting 44-47% of the spectrum (see False_Exclusion). At k = 50% every model keeps it, and then the statement ranking is unchanged because the ties are inside the buggy file.",
        "5. All 11 SFL formulas give the same ranking except Wong2: lines with identical (ef, ep, nf, np) cannot be separated by any formula.",
        "6. Breaking exact ties with the bug-report text of each statement (tie_text) is what helps: expected rank 3.5 -> 2 on Lang-1 and Lang-3, unchanged on Lang-2. Naive linear blending is worse than plain Ochiai.",
        "7. File-level score as tie-breaker (tie_file) changes nothing: tied statements share one file.",
    ]
    for i, t in enumerate(finds, start=r + 1):
        ws.cell(row=i, column=1, value=t).font = F
    r2 = r + len(finds) + 2
    ws.cell(row=r2, column=1, value="Caveats").font = FS
    cav = ["n = 3 bugs, 1 project: the differences are small counts and the bootstrap/Wilcoxon columns are not informative.",
           "tie_text / tie_file_text look best among the 12 hybrid variants on the same three bugs; this is a post-hoc comparison, so treat it as a hypothesis for more bugs.",
           "Tie-breaking by text can also move the best-case rank (Lang-1 best-case 1 -> 2); it narrows the uncertainty rather than always improving it.",
           "Statement text scores are an exploratory extension of ours, not part of the ISBL paper; the BugLocator history component is not implemented."]
    for i, t in enumerate(cav, start=r2 + 1):
        ws.cell(row=i, column=1, value=t).font = F
    ws.column_dimensions["A"].width = 66
    for c in range(2, 11):
        ws.column_dimensions[get_column_letter(c)].width = 17
    ws.column_dimensions["B"].width = 26
    ws.sheet_view.showGridLines = False
    return ws


def hybrid_modes(wb, st):
    ws = wb.create_sheet("Hybrid_Modes")
    ws["A1"] = "Hybrid variants averaged over all IR configurations and both SFL formulas (formulas over Statement_Summary)"
    ws["A1"].font = FS
    heads = ["hybrid mode", "k", "variants averaged", "mean MRR (avg ties)", "mean TOP@3 (worst ties)", "mean EXAM", "total wins vs Ochiai", "total draws", "total losses"]
    for j, h in enumerate(heads):
        ws.cell(row=2, column=1 + j, value=h)
    BE.style_header(ws, 2, 1, len(heads))
    ws.row_dimensions[2].height = 42
    n = len(st) + 1
    c = {k: col_letter(st, k) for k in ["family", "k", "hybrid", "MRR_avg", "TOP@3_worst", "EXAM_avg", "vs_SFL:ochiai_win", "vs_SFL:ochiai_draw", "vs_SFL:ochiai_lose"]}
    rng = lambda k: "Statement_Summary!$%s$2:$%s$%d" % (c[k], c[k], n)
    modes = ["tie_file", "tie_text", "tie_file_text", "linear a=0.3", "linear a=0.3 b=0.3", "mult a=0.5"]
    row = 3
    for kk in (0.5, 1.0):
        for m in modes:
            ws.cell(row=row, column=1, value=m).font = FB
            ws.cell(row=row, column=2, value=kk).font = F
            crit = '%s,"HYBRID",%s,$A%d,%s,$B%d' % (rng("family"), rng("hybrid"), row, rng("k"), row)
            ws.cell(row=row, column=3, value="=COUNTIFS(%s)" % crit)
            ws.cell(row=row, column=4, value="=AVERAGEIFS(%s,%s)" % (rng("MRR_avg"), crit)).number_format = "0.000"
            ws.cell(row=row, column=5, value="=AVERAGEIFS(%s,%s)" % (rng("TOP@3_worst"), crit)).number_format = "0.00"
            ws.cell(row=row, column=6, value="=AVERAGEIFS(%s,%s)" % (rng("EXAM_avg"), crit)).number_format = "0.0000"
            for j, k in enumerate(["vs_SFL:ochiai_win", "vs_SFL:ochiai_draw", "vs_SFL:ochiai_lose"], start=7):
                ws.cell(row=row, column=j, value="=SUMIFS(%s,%s)" % (rng(k), crit))
            for j in range(1, 10):
                ws.cell(row=row, column=j).border = BE.BOX
                if j > 2:
                    ws.cell(row=row, column=j).font = F
            row += 1
    ws.cell(row=row + 1, column=1, value="k = 1.0 means no file filter (spectrum untouched); k = 0.5 keeps the top half of the IR ranking. Plain Ochiai has mean MRR 0.302 on these bugs.").font = FI
    ws.column_dimensions["A"].width = 24
    for j in range(2, 10):
        ws.column_dimensions[get_column_letter(j)].width = 18
    ch = BarChart()
    ch.type = "col"
    ch.title = "Mean MRR (avg ties) by hybrid mode, k = 0.5"
    ch.add_data(Reference(ws, min_col=4, min_row=2, max_row=8), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=3, max_row=8))
    ch.legend = None
    ch.height, ch.width = 8.5, 18
    ws.add_chart(ch, "K2")
    return ws


def cosine_sheet(wb):
    ws = wb.create_sheet("Cosine_Audit")
    ws["A1"] = "Cosine-similarity audit (paper formula 1, raw term frequency, paper preprocessing)"
    ws["A1"].font = FT
    r = 3
    for title, name in (("1. Our cosine vs scikit-learn and SciPy", "cosine_audit_correctness.csv"),
                        ("2. Mathematical properties on the real vectors", "cosine_audit_properties.csv"),
                        ("3. Why the files score as they do: top terms and how much of the score comes from words found in >50% of files", "cosine_audit_why.csv")):
        df = csv(name)
        ws.cell(row=r, column=1, value=title).font = FS
        BE.write_table(ws, df, r0=r + 1)
        ws.row_dimensions[r + 1].height = 42
        r += len(df) + 4
    BE.autowidth(ws, maxw=60)
    ws.column_dimensions["A"].width = 14


def main():
    ir = csv("file_level_summary.csv")
    st = csv("statement_level_summary.csv")
    wb = Workbook()
    wb.remove(wb.active)
    ws = data_sheet(wb, "IR_File_Summary", ir.sort_values("MRR", ascending=False).reset_index(drop=True))
    ch = BarChart()
    ch.type = "bar"
    ch.title = "File-level MRR by IR configuration"
    ch.add_data(Reference(ws, min_col=list(ir.columns).index("MRR") + 1, min_row=1, max_row=len(ir) + 1), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=len(ir) + 1))
    ch.legend = None
    ch.height, ch.width = 14, 20
    ws.add_chart(ch, "A%d" % (len(ir) + 4))
    data_sheet(wb, "IR_File_PerBug", csv("file_level_per_bug.csv"), freeze="C2")
    data_sheet(wb, "False_Exclusion", csv("false_exclusion.csv"), freeze="A2")
    data_sheet(wb, "Statement_Summary", st, freeze="B2")
    data_sheet(wb, "Statement_PerBug", csv("statement_level_per_bug.csv"), freeze="C2")
    # keep Summary row order from the CSV for the formulas (the data sheets above are sorted only for display)
    summary(wb, ir, st)
    hybrid_modes(wb, st)
    bs = wb.create_sheet("Bugs")
    BE.write_table(bs, bugs_table())
    BE.autowidth(bs, maxw=70)
    bs.row_dimensions[1].height = 42
    cosine_sheet(wb)
    ls = wb.create_sheet("Literature")
    BE.write_table(ls, pd.DataFrame(LIT, columns=["source", "link", "what it reports", "what was added here"]))
    for i in range(2, len(LIT) + 2):
        ls.cell(row=i, column=2).hyperlink = ls.cell(row=i, column=2).value
        ls.cell(row=i, column=2).font = Font(name=BE.FONT, size=10, color="0563C1", underline="single")
    BE.autowidth(ls, maxw=80)
    gs = wb.create_sheet("Metric_Guide")
    BE.write_table(gs, pd.DataFrame(GUIDE, columns=["metric / option", "meaning"]))
    BE.autowidth(gs, maxw=110)
    order = ["Summary", "Hybrid_Modes", "IR_File_Summary", "IR_File_PerBug", "False_Exclusion", "Statement_Summary", "Statement_PerBug", "Bugs",
             "Cosine_Audit", "Literature", "Metric_Guide"]
    wb._sheets.sort(key=lambda s: order.index(s.title))
    wb.active = 0
    wb.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
