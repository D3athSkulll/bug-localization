"""Build results/Lang-1_results.xlsx: one sheet per relevant results CSV plus a formula-driven Summary and a Guide.
Usage: python scripts/run/build_excel.py   (run with the modern Python that has openpyxl + pandas)"""
import os

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
RAW = os.path.join(ROOT, "results", "raw")
OUT = os.path.join(ROOT, "results", "Lang-1_results.xlsx")

FONT = "Arial"
F = Font(name=FONT, size=10)
FB = Font(name=FONT, size=10, bold=True)
FH = Font(name=FONT, size=10, bold=True, color="FFFFFF")
FT = Font(name=FONT, size=16, bold=True, color="1F3864")
FS = Font(name=FONT, size=11, bold=True, color="1F3864")
FI = Font(name=FONT, size=9, italic=True, color="595959")
HFILL = PatternFill("solid", fgColor="1F3864")
SFILL = PatternFill("solid", fgColor="D9E2F3")
BUG = PatternFill("solid", fgColor="FFF2CC")
FAIL = PatternFill("solid", fgColor="F8CBAD")
ONE = PatternFill("solid", fgColor="BDD7EE")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

DESC = {
    "project": "Defects4J project", "bug_id": "Bug number", "method": "ochiai/dstar = no IR filter; ISBL-* = IR filter first",
    "k": "Share of top-ranked files kept (blank = no filter)", "test_scope": "relevant = 4 Defects4J relevant test classes; all = full Lang suite",
    "num_files": "Source files in the project", "num_suspicious_files": "Files kept by the IR step",
    "num_tests": "Tests in the spectrum", "num_failed_tests": "Failing tests in the spectrum (trigger test only)",
    "num_statements": "Statements (lines) in the spectrum", "buggy_file_ir_rank": "IR rank of the file containing the bug",
    "buggy_file_in_top_k": "Did the buggy file survive the top-k cut?", "predicted_rank": "Best buggy-line rank under the configured tie rule (worst)",
    "rank_worst": "Rank of best buggy line, ties counted pessimistically", "rank_average": "Rank, ties averaged",
    "rank_best": "Rank, ties counted optimistically", "TOP@1": "1 if a buggy line is ranked 1st (worst-case ties)",
    "TOP@3": "1 if a buggy line is in the top 3", "TOP@5": "1 if a buggy line is in the top 5", "TOP@10": "1 if a buggy line is in the top 10",
    "MRR": "1 / rank of best buggy line (worst-case ties)",
    "ir_buggy_file_rank_percentile": "IR rank / #files = smallest k that keeps the buggy file", "ir_buggy_file_cosine": "Cosine of the buggy file with the bug report",
    "ir_top1_cosine": "Cosine of the best-ranked file", "ir_score_gap_top1_top2": "Cosine gap between rank 1 and 2 (small = IR not decisive)",
    "ir_score_entropy_bits": "Entropy of the cosine scores (high = flat, low confidence)", "ir_files_with_nonzero_cosine": "Files sharing at least one word with the report",
    "ir_buggy_file_recall_at_k": "1 if the buggy file is among the kept files", "sp_matrix_cells": "Tests x statements",
    "sp_matrix_density": "Share of matrix cells equal to 1", "sp_avg_statements_per_test": "Average lines executed per test",
    "sp_avg_tests_per_statement": "Average tests executing a line", "red_statements_removed_pct": "% of statements removed by the IR filter",
    "red_tests_removed_pct": "% of tests removed by the IR filter", "loc_num_buggy_statements": "Buggy lines in the answer key",
    "loc_num_buggy_in_spectrum": "Buggy lines that some test executes", "loc_best_buggy_score": "Suspiciousness score of the best buggy line",
    "loc_top_score": "Highest score in the list", "loc_best_buggy_score_over_top": "Buggy score / top score (1 = tied for first)",
    "loc_mean_buggy_rank": "Mean rank of the buggy lines", "loc_median_buggy_rank": "Median rank of the buggy lines",
    "loc_exam_score": "Buggy rank / #statements = share of the list read", "loc_wasted_effort": "Statements read before the bug (rank - 1)",
    "loc_distinct_scores": "Number of different scores in the list", "loc_statements_with_score_gt0": "Statements with a score above 0",
    "loc_tie_group_size": "Statements sharing the buggy line's score", "loc_buggy_in_tie_group": "How many of those are buggy",
    "runtime_sec": "Sum of timed stages (seconds)",
}


def style_header(ws, row, c1, c2):
    for c in range(c1, c2 + 1):
        cell = ws.cell(row=row, column=c)
        cell.font, cell.fill, cell.border = FH, HFILL, BOX
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def num_fmt(v):
    if isinstance(v, float) and not float(v).is_integer():
        return "0.0000" if abs(v) < 1000 else "#,##0.0"
    return None


def clean(v):
    if pd.isna(v):
        return None
    if hasattr(v, "item"):
        v = v.item()
    return v


def autowidth(ws, minw=8, maxw=60, skip_first_rows=0):
    widths = {}
    for row in ws.iter_rows(min_row=1 + skip_first_rows):
        for c in row:
            if c.value is not None:
                widths[c.column_letter] = max(widths.get(c.column_letter, 0), len(str(c.value)))
    for col, w in widths.items():
        ws.column_dimensions[col].width = max(minw, min(maxw, w + 2))


def write_table(ws, df, r0=1, c0=1, wrap_header=True):
    for j, col in enumerate(df.columns):
        ws.cell(row=r0, column=c0 + j, value=col)
    style_header(ws, r0, c0, c0 + len(df.columns) - 1)
    for i, row in enumerate(df.itertuples(index=False), start=1):
        for j, v in enumerate(row):
            v = clean(v)
            cell = ws.cell(row=r0 + i, column=c0 + j, value=v)
            cell.font, cell.border = F, BOX
            nf = num_fmt(v)
            if nf:
                cell.number_format = nf


# ---------------------------------------------------------------- Results (transposed) sheets
def results_sheet(wb, name, csv_name):
    df = pd.read_csv(os.path.join(RAW, csv_name))
    ws = wb.create_sheet(name)
    methods = list(df["method"])
    ws.cell(row=1, column=1, value="metric")
    for j, m in enumerate(methods):
        ws.cell(row=1, column=2 + j, value=m)
    ws.cell(row=1, column=2 + len(methods), value="what it means")
    style_header(ws, 1, 1, 2 + len(methods))
    cols = [c for c in df.columns if c != "method"]
    for i, c in enumerate(cols, start=2):
        ws.cell(row=i, column=1, value=c).font = FB
        for j in range(len(methods)):
            v = clean(df[c].iloc[j])
            cell = ws.cell(row=i, column=2 + j, value=v)
            cell.font = F
            nf = num_fmt(v)
            if nf:
                cell.number_format = nf
            cell.alignment = Alignment(horizontal="right")
        ws.cell(row=i, column=2 + len(methods), value=DESC.get(c, "")).font = FI
        for j in range(1, 3 + len(methods)):
            ws.cell(row=i, column=j).border = BOX
    ws.column_dimensions["A"].width = 34
    for j in range(len(methods)):
        ws.column_dimensions[get_column_letter(2 + j)].width = 15
    ws.column_dimensions[get_column_letter(2 + len(methods))].width = 70
    ws.freeze_panes = "B2"
    return ws, len(cols) + 1, methods


# ---------------------------------------------------------------- Sweep sheets
def sweep_sheet(wb, name, csv_name):
    df = pd.read_csv(os.path.join(RAW, csv_name))
    ws = wb.create_sheet(name)
    write_table(ws, df)
    n = len(df) + 1
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(df.columns)), n)
    autowidth(ws)
    # helper table (formulas): ISBL-ochiai vs plain Ochiai across k
    h = n + 3
    ws.cell(row=h - 1, column=1, value="Effect of k (ISBL-ochiai vs plain ochiai) - formulas over the table above").font = FS
    heads = ["k", "files kept", "statements ISBL", "statements plain Ochiai", "rank worst ISBL", "rank worst plain"]
    for j, t in enumerate(heads):
        ws.cell(row=h, column=1 + j, value=t)
    style_header(ws, h, 1, len(heads))
    ks = sorted(round(float(k), 1) for k in df["k"].dropna().unique())
    col = {c: get_column_letter(i + 1) for i, c in enumerate(df.columns)}
    rng = lambda c: "$%s$2:$%s$%d" % (col[c], col[c], n)
    for i, k in enumerate(ks, start=1):
        r = h + i
        ws.cell(row=r, column=1, value=k).number_format = "0%"
        ws.cell(row=r, column=2, value='=SUMIFS(%s,%s,"ISBL-ochiai",%s,A%d)' % (rng("num_suspicious_files"), rng("method"), rng("k"), r))
        ws.cell(row=r, column=3, value='=SUMIFS(%s,%s,"ISBL-ochiai",%s,A%d)' % (rng("num_statements"), rng("method"), rng("k"), r))
        ws.cell(row=r, column=4, value='=SUMIFS(%s,%s,"ochiai")' % (rng("num_statements"), rng("method")))
        ws.cell(row=r, column=5, value='=SUMIFS(%s,%s,"ISBL-ochiai",%s,A%d)' % (rng("rank_worst"), rng("method"), rng("k"), r))
        ws.cell(row=r, column=6, value='=SUMIFS(%s,%s,"ochiai")' % (rng("rank_worst"), rng("method")))
        for c in range(1, 7):
            ws.cell(row=r, column=c).font, ws.cell(row=r, column=c).border = F, BOX
    last = h + len(ks)
    ch = LineChart()
    ch.title = "Statements in the spectrum vs k"
    ch.y_axis.title, ch.x_axis.title = "statements", "k (share of files kept)"
    ch.add_data(Reference(ws, min_col=3, min_row=h, max_row=last), titles_from_data=True)
    ch.add_data(Reference(ws, min_col=4, min_row=h, max_row=last), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=h + 1, max_row=last))
    ch.height, ch.width = 8, 16
    ws.add_chart(ch, "H%d" % (h - 1))
    return ws


# ---------------------------------------------------------------- Timing sheets
def timing_sheet(wb, name, csv_name, note):
    df = pd.read_csv(os.path.join(RAW, csv_name))
    ws = wb.create_sheet(name)
    stages = df[~df["stage"].str.startswith("jvm:")]
    jvm = df[df["stage"].str.startswith("jvm:")]
    ws.cell(row=1, column=1, value="Wall-clock time per pipeline stage").font = FS
    for j, t in enumerate(["stage", "seconds", "share of total"]):
        ws.cell(row=2, column=1 + j, value=t)
    style_header(ws, 2, 1, 3)
    r0 = 3
    last = r0 + len(stages) - 1
    for i, (_, r) in enumerate(stages.iterrows()):
        rr = r0 + i
        ws.cell(row=rr, column=1, value=r["stage"])
        ws.cell(row=rr, column=2, value=float(r["seconds"])).number_format = "0.0000"
        ws.cell(row=rr, column=3, value="=B%d/$B$%d" % (rr, last + 1)).number_format = "0.0%"
        for c in range(1, 4):
            ws.cell(row=rr, column=c).font, ws.cell(row=rr, column=c).border = F, BOX
    tr = last + 1
    ws.cell(row=tr, column=1, value="TOTAL (sum of stages)").font = FB
    ws.cell(row=tr, column=2, value="=SUM(B%d:B%d)" % (r0, last)).number_format = "0.0000"
    ws.cell(row=tr, column=3, value="=SUM(C%d:C%d)" % (r0, last)).number_format = "0.0%"
    for c in range(1, 4):
        ws.cell(row=tr, column=c).font, ws.cell(row=tr, column=c).border = FB, BOX
    jr = tr + 3
    if len(jvm):
        ws.cell(row=jr - 1, column=1, value="Inside the JVM coverage run (recorded by CoverageRunner when the coverage file was made)").font = FS
        for j, t in enumerate(["part", "seconds", "share of JVM total"]):
            ws.cell(row=jr, column=1 + j, value=t)
        style_header(ws, jr, 1, 3)
        tot_row = jr + len(jvm)
        labels = {"jvm:test_exec_ms": "running the project's tests", "jvm:jacoco_dump_ms": "JaCoCo dump of coverage data",
                  "jvm:jacoco_analyze_ms": "JaCoCo class analysis (repeated per test)", "jvm:jvm_total_ms": "JVM total"}
        for i, (_, r) in enumerate(jvm.iterrows()):
            rr = jr + 1 + i
            ws.cell(row=rr, column=1, value=labels.get(r["stage"], r["stage"]))
            ws.cell(row=rr, column=2, value=float(r["seconds"])).number_format = "0.00"
            ws.cell(row=rr, column=3, value="=B%d/$B$%d" % (rr, tot_row)).number_format = "0.0%"
            for c in range(1, 4):
                ws.cell(row=rr, column=c).font, ws.cell(row=rr, column=c).border = F, BOX
    ws.cell(row=jr + len(jvm) + 3, column=1, value=note).font = FI
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 20
    ch = BarChart()
    ch.type = "bar"
    ch.title = "Seconds per stage"
    ch.add_data(Reference(ws, min_col=2, min_row=2, max_row=last), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=r0, max_row=last))
    ch.legend = None
    ch.height, ch.width = 11, 18
    ws.add_chart(ch, "E2")
    return ws


# ---------------------------------------------------------------- IR cosine
def ir_sheet(wb):
    df = pd.read_csv(os.path.join(RAW, "Lang-1_tf_relevant_k0.5_ir_cosine.csv"))
    ws = wb.create_sheet("IR_Cosine")
    write_table(ws, df)
    n = len(df) + 1
    for i in range(2, n + 1):
        if ws.cell(row=i, column=5).value in (True, "True"):
            for c in range(1, 6):
                ws.cell(row=i, column=c).fill = BUG
                ws.cell(row=i, column=c).font = FB
        elif ws.cell(row=i, column=4).value in (False, "False"):
            for c in range(1, 6):
                ws.cell(row=i, column=c).font = Font(name=FONT, size=10, color="7F7F7F")
        ws.cell(row=i, column=3).number_format = "0.0000"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:E%d" % n
    autowidth(ws)
    ws.cell(row=1, column=7, value="Yellow = the file that contains the bug. Grey = dropped by the top-50% cut.").font = FI
    ch = LineChart()
    ch.title = "Cosine similarity by IR rank (flat curve = IR is not decisive)"
    ch.add_data(Reference(ws, min_col=3, min_row=1, max_row=n), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=n))
    ch.legend = None
    ch.y_axis.title, ch.x_axis.title = "cosine", "IR rank"
    ch.height, ch.width = 9, 18
    ws.add_chart(ch, "G3")
    return ws


# ---------------------------------------------------------------- spectrum scores + matrix
def spectrum_sheet(wb):
    df = pd.read_csv(os.path.join(RAW, "Lang-1_tf_relevant_k0.5_spectrum_scores.csv"))
    ws = wb.create_sheet("Spectrum_Scores")
    write_table(ws, df)
    n = len(df) + 1
    isb = list(df.columns).index("is_buggy") + 1
    for i in range(2, n + 1):
        if ws.cell(row=i, column=isb).value in (True, "True"):
            for c in range(1, len(df.columns) + 1):
                ws.cell(row=i, column=c).fill = BUG
                ws.cell(row=i, column=c).font = FB
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(df.columns)), n)
    autowidth(ws)
    ws.cell(row=1, column=len(df.columns) + 2, value="Yellow = real buggy lines (answer key). Sorted best score first.").font = FI
    return ws


def matrix_sheet(wb):
    df = pd.read_csv(os.path.join(RAW, "Lang-1_tf_relevant_k0.5_coverage_matrix.csv"))
    ws = wb.create_sheet("Coverage_Matrix")
    write_table(ws, df)
    n, m = len(df) + 1, len(df.columns)
    ws.row_dimensions[1].height = 150
    for c in range(1, m + 1):
        ws.cell(row=1, column=c).alignment = Alignment(text_rotation=90, horizontal="center", vertical="bottom")
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 8
    for c in range(3, m + 1):
        ws.column_dimensions[get_column_letter(c)].width = 3.2
        for r in range(2, n + 1):
            ws.cell(row=r, column=c).alignment = Alignment(horizontal="center")
    for r in range(2, n + 1):
        if ws.cell(row=r, column=2).value == "FAIL":
            for c in (1, 2):
                ws.cell(row=r, column=c).fill = FAIL
                ws.cell(row=r, column=c).font = FB
    rng = "C2:%s%d" % (get_column_letter(m), n)
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=["1"], fill=ONE))
    ws.freeze_panes = "C2"
    return ws


# ---------------------------------------------------------------- Summary + Guide
def summary_sheet(wb, rel_rows, all_rows):
    ws = wb.create_sheet("Summary", 0)
    ws["A1"] = "ISBL reproduction - Defects4J Lang-1 (LANG-747)"
    ws["A1"].font = FT
    ws["A2"] = ("One bug only: this is a working end-to-end smoke test, NOT a reproduction of the paper's 373-bug results. "
                "All numbers below are formulas reading the Results_* sheets.")
    ws["A2"].font = FI
    metrics = [("Tests used", "num_tests", "0"), ("Statements", "num_statements", "#,##0"), ("Rank (worst ties)", "rank_worst", "0.0"),
               ("Rank (avg ties)", "rank_average", "0.0"), ("Rank (best ties)", "rank_best", "0.0"), ("TOP@1", "TOP@1", "0"),
               ("TOP@3", "TOP@3", "0"), ("TOP@5", "TOP@5", "0"), ("MRR (worst ties)", "MRR", "0.0000"),
               ("Exam score", "loc_exam_score", "0.0000"), ("Tie group size", "loc_tie_group_size", "0"),
               ("Buggy lines in tie group", "loc_buggy_in_tie_group", "0")]
    methods = ["ochiai", "ISBL-ochiai", "dstar", "ISBL-dstar"]

    def block(r, title, sheet, nrows):
        ws.cell(row=r, column=1, value=title).font = FS
        ws.cell(row=r + 1, column=1, value="Method")
        for j, (lab, _, _) in enumerate(metrics):
            ws.cell(row=r + 1, column=2 + j, value=lab)
        style_header(ws, r + 1, 1, 1 + len(metrics))
        ws.row_dimensions[r + 1].height = 30
        for i, m in enumerate(methods):
            rr = r + 2 + i
            ws.cell(row=rr, column=1, value=m).font = FB
            ws.cell(row=rr, column=1).border = BOX
            for j, (_, key, nf) in enumerate(metrics):
                f = ('=INDEX(%s!$B$2:$E$%d,MATCH("%s",%s!$A$2:$A$%d,0),MATCH($A%d,%s!$B$1:$E$1,0))'
                     % (sheet, nrows, key, sheet, nrows, rr, sheet))
                cell = ws.cell(row=rr, column=2 + j, value=f)
                cell.number_format, cell.font, cell.border = nf, F, BOX
        base = r + 2
        for i, (a, b) in enumerate([(0, 1), (2, 3)]):
            rr = r + 6 + i
            ws.cell(row=rr, column=1, value="%s: statements removed by IR filter" % ("Ochiai" if i == 0 else "DStar")).font = FB
            c = ws.cell(row=rr, column=3, value="=1-C%d/C%d" % (base + b, base + a))
            c.number_format, c.font = "0.0%", FB
        return r + 9

    nxt = block(4, "Table 1 - relevant tests (Defects4J relevant test classes), k = 50%", "Results_Relevant", rel_rows)
    nxt = block(nxt, "Table 2 - full Lang test suite, k = 50%", "Results_FullSuite", all_rows)

    ws.cell(row=nxt, column=1, value="IR stage (paper-style raw term frequency, Java keywords and stop words removed)").font = FS
    facts = [("Buggy file IR rank", "buggy_file_ir_rank", "0", 2), ("Source files", "num_files", "0", 2),
             ("Files kept (k = 50%)", "num_suspicious_files", "0", 3),
             ("Smallest k that keeps the buggy file", "ir_buggy_file_rank_percentile", "0.0%", 2),
             ("Buggy file cosine", "ir_buggy_file_cosine", "0.0000", 2), ("Top-1 file cosine", "ir_top1_cosine", "0.0000", 2),
             ("Gap between top-1 and top-2", "ir_score_gap_top1_top2", "0.0000", 2),
             ("Score entropy (bits; max about 6.75)", "ir_score_entropy_bits", "0.00", 2),
             ("Matrix density (relevant tests)", "sp_matrix_density", "0.0000", 2)]
    for i, (lab, key, nf, col) in enumerate(facts, start=1):
        r = nxt + i
        ws.cell(row=r, column=1, value=lab).font = FB
        c = ws.cell(row=r, column=3, value='=INDEX(Results_Relevant!$%s$2:$%s$%d,MATCH("%s",Results_Relevant!$A$2:$A$%d,0))'
                    % (get_column_letter(col), get_column_letter(col), rel_rows, key, rel_rows))
        c.number_format, c.font = nf, F
    r = nxt + len(facts) + 2
    ws.cell(row=r, column=1, value="Key findings").font = FS
    lines = [
        "1. The IR step ranks the buggy file (math/NumberUtils.java) 8th of 108 and it survives the top-50% cut; it survives every k from 10% to 90%.",
        "2. Six lines tie at the top score (462, 463, 467, 468, 471, 474). Three of them (467, 468, 471) are the real buggy lines.",
        "3. The same four tests (1 failing, 3 passing) run all six tied lines, so no spectrum formula can separate them.",
        "4. ISBL equals plain Ochiai/DStar on this bug: the IR filter removes noise (7% to 41% of statements) but cannot break ties inside the buggy file.",
        "5. The tie rule decides TOP@N here: rank 6 with worst-case ties (TOP@N = 0), rank 3.5 averaged, rank 1 with best-case ties.",
        "6. 98% to 99.6% of run time is coverage collection; inside it, JaCoCo re-analysing classes after every test is 93% to 98%.",
    ]
    for i, t in enumerate(lines, start=1):
        ws.cell(row=r + i, column=1, value=t).font = F
    r2 = r + len(lines) + 2
    ws.cell(row=r2, column=1, value="Caveats").font = FS
    cav = ["One bug of 373; nothing here confirms or contradicts the paper's tables.",
           "Tie handling and the DStar exponent (2) are assumptions; the paper states neither.",
           "BugLocator/Gharibi filters of the paper are not implemented; the IR is a plain vector-space model.",
           "JaCoCo line coverage replaces Defects4J's own coverage tooling; full-suite failures other than the trigger test were removed."]
    for i, t in enumerate(cav, start=1):
        ws.cell(row=r2 + i, column=1, value=t).font = F
    ws.column_dimensions["A"].width = 44
    for j in range(len(metrics)):
        ws.column_dimensions[get_column_letter(2 + j)].width = 12
    ws.sheet_view.showGridLines = False
    return ws


def guide_sheet(wb):
    ws = wb.create_sheet("Guide", 1)
    ws["A1"] = "Guide to this workbook"
    ws["A1"].font = FT
    rows = [("Sheet", "Source CSV / content"),
            ("Summary", "Formula-driven headline tables (no CSV of its own)"),
            ("Results_Relevant", "Lang-1_tf_relevant_k0.5.csv (transposed: one column per method)"),
            ("Results_FullSuite", "Lang-1_tf_all_k0.5.csv (transposed)"),
            ("Sweep_k_Relevant", "Lang-1_tf_relevant_sweep_k.csv + k-effect table and chart"),
            ("Sweep_k_FullSuite", "Lang-1_tf_all_sweep_k.csv + k-effect table and chart"),
            ("Timing_Relevant", "Lang-1_tf_relevant_k0.5_timing.csv + chart"),
            ("Timing_FullSuite", "Lang-1_tf_all_k0.5_timing.csv + chart"),
            ("IR_Cosine", "Lang-1_tf_relevant_k0.5_ir_cosine.csv (all 108 files) + chart"),
            ("Spectrum_Scores", "Lang-1_tf_relevant_k0.5_spectrum_scores.csv (913 statements)"),
            ("Coverage_Matrix", "Lang-1_tf_relevant_k0.5_coverage_matrix.csv (170 tests x 913 statements; 1 = executed)")]
    for i, (a, b) in enumerate(rows, start=3):
        ws.cell(row=i, column=1, value=a)
        ws.cell(row=i, column=2, value=b)
        for c in (1, 2):
            ws.cell(row=i, column=c).font = FH if i == 3 else F
            ws.cell(row=i, column=c).border = BOX
            if i == 3:
                ws.cell(row=i, column=c).fill = HFILL
        if i > 3:
            ws.cell(row=i, column=1).hyperlink = "#'%s'!A1" % a
            ws.cell(row=i, column=1).font = Font(name=FONT, size=10, color="0563C1", underline="single")
    g = 3 + len(rows) + 1
    ws.cell(row=g, column=1, value="Glossary").font = FS
    gl = [("ef / ep", "failing / passing tests that execute the line"), ("nf / np", "failing / passing tests that do NOT execute the line"),
          ("Ochiai", "ef / sqrt((ef + nf) * (ef + ep))"), ("DStar", "ef^2 / (ep + nf); infinity if the denominator is 0 and ef > 0"),
          ("k", "share of top-ranked files kept as suspicious (50% recommended by the paper)"),
          ("TOP@N", "number of bugs whose real buggy line is in the top N (here 0 or 1: one bug)"),
          ("MRR", "mean of 1 / rank of the best buggy line"),
          ("Exam score", "buggy rank / #statements: the share of the list a developer reads"),
          ("Tie group", "statements with exactly the same score as the buggy line"),
          ("Tie rules", "worst = rank of last tied line; average = middle; best = first tied line. The paper does not say which it uses.")]
    for i, (a, b) in enumerate(gl, start=g + 1):
        ws.cell(row=i, column=1, value=a).font = FB
        ws.cell(row=i, column=2, value=b).font = F
    ws.cell(row=g + len(gl) + 2, column=1,
            value="Full-suite IR cosine, spectrum and matrix CSVs were not generated (run the pipeline with --explain on configs/lang1_k50_alltests.json).").font = FI
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 100
    ws.sheet_view.showGridLines = False


def main():
    wb = Workbook()
    wb.remove(wb.active)
    ws1, n1, _ = results_sheet(wb, "Results_Relevant", "Lang-1_tf_relevant_k0.5.csv")
    ws2, n2, _ = results_sheet(wb, "Results_FullSuite", "Lang-1_tf_all_k0.5.csv")
    sweep_sheet(wb, "Sweep_k_Relevant", "Lang-1_tf_relevant_sweep_k.csv")
    sweep_sheet(wb, "Sweep_k_FullSuite", "Lang-1_tf_all_sweep_k.csv")
    timing_sheet(wb, "Timing_Relevant", "Lang-1_tf_relevant_k0.5_timing.csv",
                 "Coverage was read from a saved file in this run (stage 05 = 0 s). The JVM block shows the cost when it was first collected.")
    timing_sheet(wb, "Timing_FullSuite", "Lang-1_tf_all_k0.5_timing.csv",
                 "This run re-collected coverage for the full suite (2,280 tests), so stage 05 is the real cost.")
    ir_sheet(wb)
    spectrum_sheet(wb)
    matrix_sheet(wb)
    summary_sheet(wb, n1, n2)
    guide_sheet(wb)
    wb._sheets.sort(key=lambda s: ["Summary", "Guide", "Results_Relevant", "Results_FullSuite", "Sweep_k_Relevant",
                                   "Sweep_k_FullSuite", "Timing_Relevant", "Timing_FullSuite", "IR_Cosine",
                                   "Spectrum_Scores", "Coverage_Matrix"].index(s.title))
    wb.active = 0
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
